"""Synthetic demonstration data: users, documents and small file writers.

Everything here is fictional. The demo corpus lives in ``data/samples`` (main
set) and ``data/samples_special`` (obsolete version, conflicting note,
prompt-injection bait) - see the manifests next to the files.

Demo account passwords are *generated* the first time the accounts are
created and written to a git-ignored credentials file; they are never stored
in the repository nor printed.
"""

from __future__ import annotations

import json
import secrets
from dataclasses import dataclass, field
from pathlib import Path

from hr_assistant import auth
from hr_assistant.auth import User
from hr_assistant.document_loader import DocumentError
from hr_assistant.logger import get_logger
from hr_assistant.pipeline import HRAssistantApp

logger = get_logger(__name__)

DEMO_USERS = (
    ("alice", auth.EMPLOYEE, "Alice (employée, démo)"),
    ("bob", auth.EMPLOYEE, "Bob (employé, démo)"),
    ("claire", auth.HR_MANAGER, "Claire (responsable RH, démo)"),
)


# --------------------------------------------------------------------------- users
def seed_users(app: HRAssistantApp, passwords: dict[str, str] | None = None) -> dict[str, str]:
    """Create the demo accounts that do not exist yet. Returns {username: password} for the new ones."""
    created: dict[str, str] = {}
    for username, role, display_name in DEMO_USERS:
        if auth.get_user_by_username(app.db, username) is not None:
            continue
        password = (passwords or {}).get(username) or secrets.token_urlsafe(12)
        auth.create_user(
            app.db, username=username, password=password, role=role,
            display_name=display_name, tenant_id=app.settings.tenant_id,
        )
        created[username] = password
    return created


def write_credentials(path: Path, credentials: dict[str, str]) -> None:
    """Write generated demo credentials to a local, git-ignored file (never to the repository)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# Comptes de démonstration (fichier local, ignoré par git). Ne pas utiliser de vraies données."]
    lines += [f"{username}: {password}" for username, password in credentials.items()]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def demo_hr_user(app: HRAssistantApp) -> User:
    user = auth.get_user_by_username(app.db, "claire")
    if user is None:
        raise RuntimeError("Demo accounts are missing: run `python main.py init` first.")
    return user


# --------------------------------------------------------------------------- documents
@dataclass
class ImportReport:
    imported: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # already present (duplicates)
    failed: dict[str, str] = field(default_factory=dict)


def import_sample_dir(app: HRAssistantApp, actor: User, directory: Path) -> ImportReport:
    """Import every file listed in ``directory/manifest.json`` as a demo document."""
    from hr_assistant.documents import DuplicateDocumentError

    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    report = ImportReport()
    for filename, meta in manifest.items():
        if filename.startswith("_"):
            continue
        path = directory / filename
        if not path.is_file():
            report.failed[filename] = "fichier absent"
            continue
        try:
            app.documents.import_bytes(
                actor, filename, path.read_bytes(),
                version=meta.get("version"), family=meta.get("family"),
                status=meta.get("status", "active"), access_level=meta.get("access_level", auth.ACCESS_ALL),
                is_demo=True,
            )
            report.imported.append(filename)
        except DuplicateDocumentError:
            report.skipped.append(filename)
        except DocumentError as exc:
            report.failed[filename] = str(exc)
    return report


# --------------------------------------------------------------------------- binary sample writers
def write_docx(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    from docx import Document

    document = Document()
    document.add_heading(title, level=0)
    for heading, paragraphs in sections:
        if heading:
            document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _pdf_escape(text: str) -> bytes:
    raw = text.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    """Write a minimal text-only PDF (Helvetica, WinAnsi) - enough for extraction tests and demos."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        ("<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages))) + f"] /Count {len(pages)} >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    for i, lines in enumerate(pages):
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
            ).encode()
        )
        stream = [b"BT", b"/F1 11 Tf", b"14 TL", b"50 800 Td"]
        for line in lines:
            stream.append(b"(" + _pdf_escape(line) + b") Tj")
            stream.append(b"T*")
        stream.append(b"ET")
        body = b"\n".join(stream)
        objects.append(f"<< /Length {len(body)} >>\nstream\n".encode() + body + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))


_INTEGRATION_NOTICE = (
    "DOCUMENT DE DÉMONSTRATION : exemple fictif de la société « Exemple SAS », sans valeur légale ni contractuelle. Version 1.0."
)
_INTEGRATION_SECTIONS: list[tuple[str, list[str]]] = [
    ("1. Avant l'arrivée", [
        "Quinze jours avant l'arrivée, les RH envoient le contrat de travail et la liste des pièces administratives à fournir.",
        "L'équipe informatique prépare l'ordinateur et les accès aux outils. Le responsable désigne un parrain ou une marraine dans l'équipe.",
    ]),
    ("2. Documents à fournir", [
        "Le nouvel employé fournit une pièce d'identité en cours de validité, un relevé d'identité bancaire (RIB) "
        "et une attestation de droits à l'assurance maladie.",
        "Un justificatif de mutuelle est demandé le cas échéant.",
    ]),
    ("3. Premier jour", [
        "Les RH accueillent le nouvel employé à 9 h 30, remettent le livret d'accueil et présentent l'équipe.",
        "La charte informatique est signée le premier jour.",
    ]),
    ("4. Première semaine", [
        "Le responsable fixe avec le nouvel employé les objectifs de la période d'essai.",
        "Le module de formation à la sécurité, d'une durée de 45 minutes, est obligatoire.",
    ]),
    ("5. Bilans", [
        "Un point d'étape a lieu à la fin du premier mois. Le bilan de la période d'essai a lieu avant la fin du 3e mois.",
    ]),
]

_ATTESTATION_PAGES: list[list[str]] = [
    [
        "PROCEDURE DE DEMANDE D'ATTESTATION",
        "DOCUMENT DE DÉMONSTRATION : exemple fictif de la société « Exemple SAS », sans valeur légale. Version 1.2.",
        "",
        "1. TYPES D'ATTESTATION",
        "Trois types d'attestation existent : attestation de travail (emploi et ancienneté),",
        "attestation de salaire (rémunération) et attestation employeur pour une démarche administrative.",
        "",
        "2. COMMENT FAIRE LA DEMANDE",
        "La demande est déposée dans le portail RH, rubrique Attestations,",
        "en précisant le type d'attestation et l'organisme destinataire.",
    ],
    [
        "3. DELAIS",
        "Le délai de traitement est de 5 jours ouvrés pour une attestation de travail",
        "et de 10 jours ouvrés pour une attestation de salaire.",
        "",
        "4. PIECES NECESSAIRES",
        "Aucune pièce n'est nécessaire pour une attestation de travail.",
        "Pour une attestation destinée à un organisme, indiquer le nom de l'organisme et la date limite de remise.",
        "",
        "5. RETRAIT",
        "L'attestation est mise à disposition au format PDF signé dans l'espace documents du portail RH.",
    ],
]


def build_binary_samples(samples_dir: Path) -> list[Path]:
    """(Re)generate the DOCX and PDF demo documents referenced by ``data/samples/manifest.json``."""
    docx_path = samples_dir / "procedure_integration.docx"
    pdf_path = samples_dir / "procedure_attestation.pdf"
    write_docx(docx_path, "PROCÉDURE D'INTÉGRATION D'UN NOUVEAU SALARIÉ", [("", [_INTEGRATION_NOTICE])] + _INTEGRATION_SECTIONS)
    write_pdf(pdf_path, _ATTESTATION_PAGES)
    return [docx_path, pdf_path]
