"""Step 1: validate an uploaded file and extract its text (TXT, PDF, DOCX).

Supported on purpose: ``.txt``, ``.pdf`` (text layer only - scanned PDFs have
no extractable text and are reported as an extraction error, OCR is not
supported) and ``.docx``. Anything else is rejected with an explicit message
instead of being silently mis-read.

PDF pages are kept (page numbers feed the citations). DOCX has no reliable
page notion, so its text is returned as a single page without a number;
heading styles are preserved as ``## `` markers so the splitter can recover
the section structure.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from hr_assistant.logger import get_logger
from hr_assistant.textutils import detect_language

logger = get_logger(__name__)

SUPPORTED_TYPES = {".txt": "txt", ".pdf": "pdf", ".docx": "docx"}
_MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class DocumentError(Exception):
    """Base class: the document cannot be imported. The message is safe to show to the user."""


class UnsupportedFormatError(DocumentError):
    pass


class ExtractionError(DocumentError):
    pass


@dataclass
class Page:
    number: int | None
    text: str


@dataclass
class ExtractedDocument:
    doc_type: str
    pages: list[Page] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def language(self) -> str:
        return detect_language(self.text)


def sanitize_filename(name: str) -> str:
    """Keep only the base name and safe characters (no path traversal, no control chars)."""
    base = Path(name.replace("\\", "/")).name
    base = re.sub(r"[^\w.\- ()]", "_", base, flags=re.UNICODE).strip(" .")
    return base or "document"


def detect_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        return SUPPORTED_TYPES[ext]
    except KeyError:
        supported = ", ".join(sorted(SUPPORTED_TYPES))
        raise UnsupportedFormatError(
            f"Format non pris en charge : '{ext or 'sans extension'}'. Formats acceptés : {supported}."
        ) from None


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Check extension, size and magic bytes. Returns the document type."""
    doc_type = detect_type(filename)
    if not data:
        raise ExtractionError("Le fichier est vide.")
    if len(data) > max_bytes:
        raise DocumentError(
            f"Fichier trop volumineux ({len(data) // (1024 * 1024)} Mo, maximum {max_bytes // (1024 * 1024)} Mo)."
        )
    if doc_type == "pdf" and not data.lstrip().startswith(b"%PDF-"):
        raise ExtractionError("Le fichier n'est pas un PDF valide (signature absente).")
    if doc_type == "docx" and not data.startswith(b"PK"):
        raise ExtractionError("Le fichier n'est pas un DOCX valide (archive absente).")
    return doc_type


def _decode_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ExtractionError("Le fichier .txt semble binaire (octets nuls détectés).")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Encodage du fichier texte non reconnu (UTF-8 ou Windows-1252 attendu).")


def _extract_pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Le PDF est protégé par un mot de passe.")
        pages = [Page(number=i, text=(page.extract_text() or "").strip()) for i, page in enumerate(reader.pages, start=1)]
    except ExtractionError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError) as exc:
        logger.warning("PDF extraction failed: %s", type(exc).__name__)
        raise ExtractionError("Le PDF est illisible ou corrompu.") from exc
    if not any(p.text for p in pages):
        raise ExtractionError(
            "Aucun texte extractible dans ce PDF (document scanné ? l'OCR n'est pas pris en charge)."
        )
    return pages


def _extract_docx(data: bytes) -> list[Page]:
    from docx import Document as DocxDocument

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > _MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ExtractionError("Le DOCX est trop volumineux une fois décompressé.")
        document = DocxDocument(io.BytesIO(data))
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            lines.append(f"## {text}" if style.lower().startswith(("heading", "title", "titre")) else text)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
    except ExtractionError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        logger.warning("DOCX extraction failed: %s", type(exc).__name__)
        raise ExtractionError("Le DOCX est illisible ou corrompu.") from exc
    text = "\n".join(lines).strip()
    if not text:
        raise ExtractionError("Le DOCX ne contient aucun texte.")
    return [Page(number=None, text=text)]


def extract_text_from_bytes(filename: str, data: bytes, max_bytes: int = 10 * 1024 * 1024) -> ExtractedDocument:
    """Validate then extract. Raises ``DocumentError`` subclasses with user-safe messages."""
    doc_type = validate_upload(filename, data, max_bytes)
    if doc_type == "txt":
        text = _decode_text(data).replace("\r\n", "\n").strip()
        if not text:
            raise ExtractionError("Le fichier est vide.")
        pages = [Page(number=None, text=text)]
    elif doc_type == "pdf":
        pages = _extract_pdf(data)
    else:
        pages = _extract_docx(data)
    logger.info("Extracted %s: %d page(s), %d char(s)", doc_type, len(pages), sum(len(p.text) for p in pages))
    return ExtractedDocument(doc_type=doc_type, pages=pages)


def load_document(file_path: str | Path, max_bytes: int = 10 * 1024 * 1024) -> ExtractedDocument:
    """Extract a file from disk (kept under the original name for backward compatibility)."""
    path = Path(file_path)
    return extract_text_from_bytes(path.name, path.read_bytes(), max_bytes)
